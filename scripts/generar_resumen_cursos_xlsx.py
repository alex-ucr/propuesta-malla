#!/usr/bin/env python3
"""
Genera resumen-cursos.xlsx con sigla, nombre, créditos, requisitos y
correquisitos de cada programa en Cursos/*-cuerpo.tex.

Uso (desde la raíz del proyecto):
    python scripts/generar_resumen_cursos_xlsx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CURSOS = ROOT / "Cursos"
OUTPUT = ROOT / "resumen-cursos.xlsx"

TITLE_RE = re.compile(r"\\(?:subsection|paragraph)\*\{([^}]*)\}")
TABULAR_RE = re.compile(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", re.DOTALL)
LABEL_RE = re.compile(r"\\textbf\{([^}:]+):\s*\}?")
SIGLA_RE = re.compile(r"^([A-Z]{2})-?(\d{4})\s+(.*)$")

ACCENTS = {
    r"\'a": "á", r"\'e": "é", r"\'i": "í", r"\'\i": "í", r"\'o": "ó", r"\'u": "ú",
    r"\'A": "Á", r"\'E": "É", r"\'I": "Í", r"\'O": "Ó", r"\'U": "Ú",
    r"\~n": "ñ", r"\~N": "Ñ", r'\"u': "ü",
}


def strip_comments(text: str) -> str:
    return "\n".join(re.sub(r"(?<!\\)%.*", "", line) for line in text.splitlines())


def clean(text: str) -> str:
    for src, dst in ACCENTS.items():
        text = text.replace("{" + src + "}", dst).replace(src, dst)
    text = re.sub(r"\{\\'\{?\\?([aeiouAEIOU])\}?\}", lambda m: ACCENTS.get("\\'" + m.group(1), m.group(1)), text)
    text = re.sub(r"\\(?:textbf|textit|emph|text)\{([^}]*)\}", r"\1", text)
    text = text.replace("~", " ").replace("\\&", "&").replace("--", "–")
    text = re.sub(r"\\[a-zA-Z]+\*?", "", text)
    text = text.replace("{", "").replace("}", "").replace("$", "")
    return re.sub(r"\s+", " ", text).strip(" ;,")


def parse_table(body: str) -> dict[str, str]:
    """Campos del cuadro de encabezado; une las filas de continuación ('\\\\ &')."""
    fields: dict[str, str] = {}
    last_label_by_col: dict[int, str] = {}
    body = body.replace("\\hline", "")
    for row in re.split(r"\\\\", body):
        if "\\multicolumn" in row:
            continue
        for col, cell in enumerate(row.split("&")):
            cell = cell.strip()
            if not cell:
                continue
            labels = list(LABEL_RE.finditer(cell))
            if labels:
                for i, m in enumerate(labels):
                    end = labels[i + 1].start() if i + 1 < len(labels) else len(cell)
                    label = clean(m.group(1))
                    fields[label] = cell[m.end():end].strip()
                    last_label_by_col[col] = label
            elif col in last_label_by_col:
                label = last_label_by_col[col]
                fields[label] = f"{fields[label]} {cell}"
    return {k: clean(v) for k, v in fields.items()}


def parse_curso(path: Path) -> dict[str, str] | None:
    text = strip_comments(path.read_text(encoding="utf-8"))
    title_match = TITLE_RE.search(text)
    table_match = TABULAR_RE.search(text)
    if not title_match or not table_match:
        return None

    title = clean(title_match.group(1))
    sigla_match = SIGLA_RE.match(title)
    if sigla_match:
        sigla = f"{sigla_match.group(1)}-{sigla_match.group(2)}"
        nombre = sigla_match.group(3)
    else:
        sigla, nombre = "", title
    if path.stem.endswith("-propuesta-cuerpo"):
        nombre = f"{nombre} (propuesta)"

    fields = parse_table(table_match.group(1))
    creditos_raw = fields.get("Créditos", "")
    creditos: int | str = int(creditos_raw) if creditos_raw.isdigit() else creditos_raw
    return {
        "Sigla": sigla,
        "Nombre": nombre,
        "Créditos": creditos,
        "Requisitos": fields.get("Requisitos", ""),
        "Correquisitos": fields.get("Correquisitos", ""),
    }


def main() -> int:
    cuerpos = sorted(CURSOS.glob("*-cuerpo.tex"))
    cursos, omitidos = [], []
    for path in cuerpos:
        curso = parse_curso(path)
        if curso is None:
            omitidos.append(path.name)
        else:
            cursos.append(curso)
    cursos.sort(key=lambda c: (c["Sigla"] or "~", c["Nombre"]))

    columns = ["Sigla", "Nombre", "Créditos", "Requisitos", "Correquisitos"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cursos"
    ws.append(columns)
    for curso in cursos:
        ws.append([curso[c] for c in columns])

    header_fill = PatternFill("solid", fgColor="005DA4")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].alignment = Alignment(horizontal="center", vertical="top")

    for i, width in enumerate([11, 48, 10, 60, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT)
    print(f"{len(cursos)} cursos escritos en {OUTPUT.name}")
    for name in omitidos:
        print(f"Omitido (sin título o cuadro de encabezado): {name}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
